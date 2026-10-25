"""
Generates the 12-slide Viva Presentation PowerPoint (.pptx) using python-pptx.
Implements 16:9 widescreen slides with Minimalist Peach executive styling,
crisp white cards, warm peach accents, high-contrast typography, and speaker notes.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

# Minimalist Peach Palette
BG_COLOR = RGBColor(253, 250, 246)       # Soft warm linen #FDFAF6
CARD_BG = RGBColor(255, 255, 255)        # Crisp clean white #FFFFFF
CARD_SUBTLE = RGBColor(255, 247, 243)    # Delicate peach tint #FFF7F3
CARD_BORDER = RGBColor(237, 227, 218)    # Hairline peach/stone #EDE3DA
PEACH_PRIMARY = RGBColor(231, 111, 81)   # Primary Peach/Coral #E76F51
PEACH_DARK = RGBColor(185, 66, 38)       # Deep terracotta #B94226
TEXT_PRIMARY = RGBColor(35, 31, 30)      # Deep warm espresso #231F1E
TEXT_MUTED = RGBColor(110, 100, 94)      # Warm slate grey #6E645E
EMERALD = RGBColor(22, 101, 52)          # Deep forest sage #166534
ROSE = RGBColor(185, 28, 28)             # Deep crimson #B91C1C
AMBER = RGBColor(154, 52, 18)            # Deep golden amber #9A3412


def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="CLASSICAL AI CYBERSECURITY"):
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = PEACH_PRIMARY

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.5), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_PRIMARY

    def add_card(slide, left, top, width, height, title=None):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1.2)

        if title:
            tb = slide.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.15), Inches(width - 0.4), Inches(0.4))
            tf = tb.text_frame
            p = tf.paragraphs[0]
            p.text = title
            p.font.size = Pt(13)
            p.font.bold = True
            p.font.color.rgb = PEACH_DARK
        return card

    def add_callout(slide, left, top, width, height, text, accent_color=PEACH_PRIMARY):
        callout = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        callout.fill.solid()
        callout.fill.fore_color.rgb = CARD_SUBTLE
        callout.line.color.rgb = accent_color
        callout.line.width = Pt(1.5)
        tf = callout.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = text
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = TEXT_PRIMARY

    def set_speaker_notes(slide, notes_text):
        notes_slide = slide.notes_slide
        tf = notes_slide.notes_text_frame
        tf.text = notes_text

    # =========================================================================
    # SLIDE 1: Title
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.2), Inches(4.0), Inches(0.45))
    badge.fill.solid()
    badge.fill.fore_color.rgb = CARD_SUBTLE
    badge.line.color.rgb = PEACH_PRIMARY
    p_b = badge.text_frame.paragraphs[0]
    p_b.text = "COLLEGE AI PROJECT :: SYMBOLIC AI"
    p_b.font.size = Pt(10)
    p_b.font.bold = True
    p_b.font.color.rgb = PEACH_DARK

    t_box = s1.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.5), Inches(1.8))
    tf_t = t_box.text_frame
    tf_t.word_wrap = True
    p1 = tf_t.paragraphs[0]
    p1.text = "Explainable URL & QR Risk Screening\nUsing Classical AI"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_PRIMARY

    p2 = tf_t.add_paragraph()
    p2.text = "Offline, Deterministic and Explainable Security Analysis"
    p2.font.size = Pt(18)
    p2.font.color.rgb = PEACH_DARK

    add_card(s1, 0.8, 4.0, 11.7, 1.4, "FOUNDATIONAL AI CONCEPTS")
    c_box = s1.shapes.add_textbox(Inches(1.0), Inches(4.5), Inches(11.3), Inches(0.8))
    p_c = c_box.text_frame.paragraphs[0]
    p_c.text = "• Propositional Definite Horn Logic     • PL-FC-Entails (Forward Chaining)\n• A* Heuristic Search (f=g+h)          • Proof-Tree Derivation DAGs          • Offline QR Decoding"
    p_c.font.size = Pt(13)
    p_c.font.bold = True
    p_c.font.color.rgb = TEXT_PRIMARY

    auth_box = s1.shapes.add_textbox(Inches(0.8), Inches(5.8), Inches(11.5), Inches(1.0))
    p_a = auth_box.text_frame.paragraphs[0]
    p_a.text = "Department of Computer Science and Engineering\nRajalakshmi Engineering College"
    p_a.font.size = Pt(13)
    p_a.font.color.rgb = TEXT_MUTED

    set_speaker_notes(
        s1,
        "Our project is an offline URL and QR risk-screening system based entirely on classical and symbolic AI. "
        "Instead of using machine learning, we represent URL characteristics as logical facts and use Horn-rule "
        "forward chaining to derive security patterns. We additionally use A* search for detecting possible brand typosquatting."
    )

    # =========================================================================
    # SLIDE 2: Problem Statement
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "Problem Statement", "CONTEXT & MOTIVATION")

    add_card(s2, 0.8, 1.6, 5.7, 4.3, "CURRENT THREAT LANDSCAPE")
    tb = s2.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(5.3), Inches(3.5))
    tf = tb.text_frame
    tf.word_wrap = True
    bullets = [
        "Suspicious URLs and phishing campaigns are widely delivered via direct links and QR codes.",
        "Visual appearance alone does not reveal deep structural attack vectors (userinfo '@' masking, IP hosts, non-standard ports).",
        "Modern detection systems rely heavily on:",
        "    • Black-box Machine Learning models",
        "    • Live external threat intelligence feeds",
        "    • Live DNS lookups and web crawlers",
        "    • Cloud API dependencies",
    ]
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = b
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_PRIMARY

    add_card(s2, 6.8, 1.6, 5.7, 4.3, "OUR CHALLENGE & FOCUS")
    tb2 = s2.shapes.add_textbox(Inches(7.0), Inches(2.2), Inches(5.3), Inches(3.5))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    bullets2 = [
        "Cloud-dependent systems fail in air-gapped, privacy-sensitive, or zero-connectivity environments.",
        "Neural networks and statistical classifiers cannot provide formal, auditable proof of their classification.",
        "Core Research Question:",
        "How can a URL or QR-encoded link be analyzed completely locally using Classical AI techniques to produce an explainable, deterministic security verdict?",
    ]
    for i, b in enumerate(bullets2):
        p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
        p.text = b
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_PRIMARY if i < 2 else PEACH_DARK

    add_callout(s2, 0.8, 6.1, 11.7, 0.8, "Focus: Deterministic structural analysis with verifiable causal proof, not probabilistic guessing.")
    set_speaker_notes(
        s2,
        "The main problem we address is not trying to detect every malicious URL in existence. Instead, we focus on "
        "deterministic structural analysis. The system should explain exactly which URL characteristics caused a security pattern to be derived."
    )

    # =========================================================================
    # SLIDE 3: Objectives
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "System Objectives", "PROJECT GOALS")

    add_card(s3, 0.8, 1.6, 5.7, 4.8, "CORE REASONING OBJECTIVES")
    tb = s3.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(5.3), Inches(4.0))
    tf = tb.text_frame
    tf.word_wrap = True
    objs1 = [
        "1. Strictly Offline Analysis: Zero network queries, DNS lookups, or external threat APIs.",
        "2. Lexical Fact Extraction: Deterministically extract ground propositional security facts.",
        "3. Horn Knowledge Base: Model expert security knowledge via pure positive definite Horn clauses.",
        "4. Forward Chaining: Implement PL-FC-Entails to infer security patterns with full provenance.",
        "5. A* Heuristic Search: Model brand typosquatting as optimal string transformation pathfinding.",
    ]
    for i, o in enumerate(objs1):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = o
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_PRIMARY

    add_card(s3, 6.8, 1.6, 5.7, 4.8, "OUTCOMES & TRANSPARENCY")
    tb2 = s3.shapes.add_textbox(Inches(7.0), Inches(2.2), Inches(5.3), Inches(4.0))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    objs2 = [
        "6. Categorical Risk Verdicts: Map inferred patterns strictly to SAFE, REVIEW, SUSPICIOUS, or HIGH RISK.",
        "7. Explainable Proof Trees: Synthesize recursive derivation DAGs answering 'Why was this verdict assigned?'.",
        "8. Offline QR Decoding: Local OpenCV 2D matrix extraction providing complete semantic parity.",
        "9. Strict Determinism: Identical inputs produce identical facts, proofs, and verdicts every time.",
    ]
    for i, o in enumerate(objs2):
        p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
        p.text = o
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_PRIMARY

    set_speaker_notes(
        s3,
        "Our objectives were not only detection, but also explainability and reproducibility. Given the same input, "
        "the system should produce the same facts, deductions and verdict every time."
    )

    # =========================================================================
    # SLIDE 4: Proposed Solution
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "Proposed Solution & AI Architecture", "METHODOLOGY")

    add_card(s4, 0.8, 1.6, 5.7, 4.5, "PROCESSING PIPELINE")
    tb = s4.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(5.3), Inches(3.7))
    tf = tb.text_frame
    tf.word_wrap = True
    steps = [
        "1. Input: Direct URL string or Offline QR Image",
        "2. Offline Input Adapter: Local matrix decoding (OpenCV)",
        "3. Lexical Parser: Strict RFC 3986 component decomposition",
        "4. Feature Extraction: Positive propositional facts & evidence",
        "5. Knowledge Base: Registry of definite Horn clauses",
        "6. Forward Chaining: PL-FC-Entails deduction",
        "7. Risk Evaluator: Max-severity categorical resolution",
        "8. Proof Tree Synthesis: Causal Derivation DAG",
    ]
    for i, st in enumerate(steps):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = st
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_PRIMARY

    add_card(s4, 6.8, 1.6, 5.7, 2.1, "PRIMARY AI ALGORITHM")
    tb_p = s4.shapes.add_textbox(Inches(7.0), Inches(2.1), Inches(5.3), Inches(1.4))
    tf_p = tb_p.text_frame
    p_p = tf_p.paragraphs[0]
    p_p.text = "PL-FC-Entails / Forward Chaining\n• Operates over declarative Horn clauses\n• Derives security patterns with complete provenance trace"
    p_p.font.size = Pt(12)
    p_p.font.color.rgb = TEXT_PRIMARY

    add_card(s4, 6.8, 3.9, 5.7, 2.2, "SECONDARY AI ALGORITHM")
    tb_s = s4.shapes.add_textbox(Inches(7.0), Inches(4.4), Inches(5.3), Inches(1.5))
    tf_s = tb_s.text_frame
    p_s = tf_s.paragraphs[0]
    p_s.text = "A* Heuristic Search\n• State-space string transformation search (f = g + h)\n• Emits ground fact: brand_typosquatting_detected"
    p_s.font.size = Pt(12)
    p_s.font.color.rgb = TEXT_PRIMARY

    add_callout(s4, 0.8, 6.3, 11.7, 0.7, "KEY PRINCIPLE: A* identifies brand similarity; Forward Chaining performs the security reasoning.")
    set_speaker_notes(
        s4,
        "Notice the fundamental separation of concerns: A* is strictly upstream feature discovery. "
        "A* does not produce a risk verdict. It discovers whether an observed label transforms into a known brand. "
        "Forward Chaining takes that fact and performs the actual cybersecurity reasoning."
    )

    # =========================================================================
    # SLIDE 5: System Architecture
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "End-to-End System Architecture", "DATA FLOW DIAGRAM")

    add_card(s5, 0.8, 1.6, 3.6, 4.4, "INPUT & PREPROCESSING")
    tb = s5.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(3.2), Inches(3.6))
    tf = tb.text_frame
    bullets_flow1 = [
        "[Input Mode]",
        "• Direct URL string",
        "• Offline QR Code image",
        "",
        "[QR Matrix Decoder]",
        "• Local OpenCV detection",
        "• Validates URL syntax",
        "",
        "[RFC 3986 URL Parser]",
        "• Scheme, Authority, Path",
        "• IP detection (IPv4 / IPv6)",
    ]
    for i, b in enumerate(bullets_flow1):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = b
        p.font.size = Pt(11)
        p.font.color.rgb = PEACH_DARK if b.startswith("[") else TEXT_PRIMARY

    add_card(s5, 4.8, 1.6, 3.7, 4.4, "SYMBOLIC AI REASONING")
    tb = s5.shapes.add_textbox(Inches(5.0), Inches(2.2), Inches(3.3), Inches(3.6))
    tf = tb.text_frame
    bullets_flow2 = [
        "[Fact Generator]",
        "• 12 base positive Horn literals",
        "",
        "[A* Typosquatting Search]",
        "• f(n) = g(n) + h(n)",
        "• Emits brand_typosquatting_detected",
        "",
        "[Knowledge Base & FC]",
        "• Declarative Horn Clauses",
        "• PL-FC-Entails Agenda Engine",
        "• Derivation Step Tracking",
    ]
    for i, b in enumerate(bullets_flow2):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = b
        p.font.size = Pt(11)
        p.font.color.rgb = PEACH_DARK if b.startswith("[") else TEXT_PRIMARY

    add_card(s5, 8.9, 1.6, 3.6, 4.4, "EVALUATION & PROOF")
    tb = s5.shapes.add_textbox(Inches(9.1), Inches(2.2), Inches(3.2), Inches(3.6))
    tf = tb.text_frame
    bullets_flow3 = [
        "[Deterministic Evaluator]",
        "• Resolves maximum severity",
        "• CRITICAL -> HIGH RISK",
        "• HIGH -> SUSPICIOUS",
        "• MEDIUM / LOW -> REVIEW",
        "• Clean Profile -> SAFE",
        "",
        "[Proof Tree Synthesis]",
        "• Recursive Derivation DAG",
        "• Concrete Lexical Evidence",
    ]
    for i, b in enumerate(bullets_flow3):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = b
        p.font.size = Pt(11)
        p.font.color.rgb = PEACH_DARK if b.startswith("[") else TEXT_PRIMARY

    add_callout(s5, 0.8, 6.2, 11.7, 0.7, "100% LOCAL & AIR-GAPPED — Zero DNS queries, HTTP calls, or external threat intelligence APIs.")
    set_speaker_notes(
        s5,
        "Here is our complete system architecture. Notice how every phase is decoupled. The QR decoder is strictly "
        "an input adapter. The URL parser decomposes the URL according to RFC 3986. The fact generator asserts ground literals. "
        "Forward Chaining derives security patterns, and the Risk Evaluator produces a deterministic categorical verdict."
    )

    # =========================================================================
    # SLIDE 6: Classical AI Techniques
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "Classical AI Techniques Breakdown", "SYMBOLIC FOUNDATIONS")

    table_shape = s6.shapes.add_table(6, 3, Inches(0.8), Inches(1.6), Inches(11.7), Inches(3.4))
    table = table_shape.table
    table.columns[0].width = Inches(2.6)
    table.columns[1].width = Inches(2.3)
    table.columns[2].width = Inches(6.8)

    headers = ["Classical AI Concept", "System Role", "Implementation in System"]
    for col_idx, h in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = CARD_SUBTLE
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = PEACH_DARK

    data = [
        ("Propositional Logic", "Ground Facts", "Encodes observed URL traits as atomic positive symbols (e.g. hostname_is_ip_address)"),
        ("Definite Horn Clauses", "Knowledge Base", "Formal cybersecurity axioms of the form P1 ^ P2 ^ ... ^ Pn => Q"),
        ("Forward Chaining", "Primary Inference", "Textbook PL-FC-Entails with FIFO agenda, count table, premise indexing"),
        ("A* Heuristic Search", "Secondary Search", "Pathfinding in discrete string edit space with admissible heuristic h(n) = |Ls - Lt|"),
        ("Proof Tree Synthesis", "Explainable AI (XAI)", "Recursive DAG traversal generating complete causal justification for every verdict"),
    ]
    for row_idx, row in enumerate(data, start=1):
        for col_idx, val in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_BG
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(11)
            p.font.color.rgb = TEXT_PRIMARY

    add_callout(s6, 0.8, 5.3, 11.7, 0.8, "STRICT CONSTRAINT: No Machine Learning, Deep Learning, Computer Vision models, or LLMs.", ROSE)
    add_callout(s6, 0.8, 6.3, 11.7, 0.6, "Primary AI: PL-FC-Entails (Forward Chaining)  |  Secondary AI: A* Heuristic Search", PEACH_PRIMARY)
    set_speaker_notes(
        s6,
        "The project is fundamentally symbolic. The URL is converted into propositions such as hostname_is_ip_address "
        "or path_contains_sensitive_keyword. Horn clauses describe relationships between these propositions, and Forward "
        "Chaining derives higher-level security patterns."
    )

    # =========================================================================
    # SLIDE 7: Knowledge Representation
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "Knowledge Representation Using Horn Clauses", "FORMAL LOGIC")

    add_card(s7, 0.8, 1.6, 5.7, 4.4, "HORN CLAUSE SPECIFICATION")
    tb = s7.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(5.3), Inches(3.6))
    tf = tb.text_frame
    tf.word_wrap = True
    body_s7 = [
        "Definite Horn Clause Form:",
        "P1 ^ P2 ^ ... ^ Pn  =>  Q",
        "",
        "• Pure positive atomic literals (negation is strictly prohibited).",
        "• Exactly one positive literal in the head (conclusion Q).",
        "• Each rule annotated with an immutable SeverityLevel:",
        "    CRITICAL, HIGH, MEDIUM, LOW, or NONE.",
        "",
        "Example Rule (RULE_IP_AND_SENSITIVE_KEYWORD):",
        "hostname_is_ip_address",
        "    ^ path_contains_sensitive_keyword",
        "    => pattern_credential_harvesting_on_ip [CRITICAL]",
    ]
    for i, b in enumerate(body_s7):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = b
        p.font.size = Pt(11)
        p.font.color.rgb = PEACH_DARK if "=>" in b or "Form:" in b else TEXT_PRIMARY

    add_card(s7, 6.8, 1.6, 5.7, 4.4, "KNOWLEDGE BASE RULE CATALOG")
    tb2 = s7.shapes.add_textbox(Inches(7.0), Inches(2.2), Inches(5.3), Inches(3.6))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    rules_cat = [
        "1. Direct IP + Sensitive Keyword [CRITICAL]",
        "2. Authority Userinfo '@' Masking [CRITICAL]",
        "3. Typosquatting + Insecure HTTP [HIGH]",
        "4. Typosquatting + Sensitive Keyword [CRITICAL]",
        "5. Redirect + Excessive Subdomains [HIGH]",
        "6. Executable Double Extension [CRITICAL]",
        "7. Abnormal URI Scheme (javascript:, data:) [CRITICAL]",
        "8. Insecure Transport on Benign Path [LOW]",
        "9. Non-Standard Network Port [MEDIUM]",
        "10. Clean Baseline Profile (11 Clean Indicators) [NONE]",
    ]
    for i, r in enumerate(rules_cat):
        p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
        p.text = r
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_PRIMARY

    add_callout(s7, 0.8, 6.2, 11.7, 0.7, "The knowledge base encodes expert cybersecurity axioms; Forward Chaining mechanically deduces their consequences.")
    set_speaker_notes(
        s7,
        "The important point is that the rules themselves contain the security knowledge. Forward Chaining doesn't know that a direct IP "
        "combined with a login path is suspicious by itself; that relationship is explicitly encoded in the knowledge base."
    )

    # =========================================================================
    # SLIDE 8: Forward Chaining
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "PL-FC-Entails (Forward Chaining)", "PRIMARY AI REASONING ENGINE")

    add_card(s8, 0.8, 1.6, 5.7, 4.4, "ALGORITHM ARCHITECTURE")
    tb = s8.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(5.3), Inches(3.6))
    tf = tb.text_frame
    tf.word_wrap = True
    fc_steps = [
        "1. Initialize Agenda queue (FIFO) with ground initial facts.",
        "2. Initialize count[rule] = len(rule.premises) for all rules.",
        "3. Pop symbol 'p' from Agenda; mark as inferred.",
        "4. For each rule 'r' indexed with premise 'p':",
        "    • Decrement count[r] by 1.",
        "    • If count[r] == 0 (all premises satisfied):",
        "        - Record DerivationStep provenance.",
        "        - Add conclusion 'q' to Agenda.",
        "5. Repeat until Agenda is empty.",
        "",
        "Guarantees: Soundness, Completeness, O(N) linear time, Monotonicity.",
    ]
    for i, st in enumerate(fc_steps):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = st
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_PRIMARY if not st.startswith("Guarantees:") else PEACH_DARK

    add_card(s8, 6.8, 1.6, 5.7, 4.4, "CONCRETE DEDUCTION TRACE")
    tb2 = s8.shapes.add_textbox(Inches(7.0), Inches(2.2), Inches(5.3), Inches(3.6))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    trace_steps = [
        "Observed URL: https://192.168.1.10/login",
        "",
        "Ground Facts Asserted:",
        "    • hostname_is_ip_address",
        "    • path_contains_sensitive_keyword",
        "",
        "Matching Rule: RULE_IP_AND_SENSITIVE_KEYWORD",
        "Premise Count: count[rule] decreases 2 -> 1 -> 0",
        "",
        "Derived Conclusion: pattern_credential_harvesting_on_ip",
        "Severity Assigned: CRITICAL",
        "Final Categorical Verdict: HIGH RISK",
    ]
    for i, t in enumerate(trace_steps):
        p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
        p.text = t
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_PRIMARY if not t.startswith("Final") else ROSE

    add_callout(s8, 0.8, 6.2, 11.7, 0.7, 'VIVA LINE: "Forward Chaining derives what logically follows from the facts and rules; it does not predict."')
    set_speaker_notes(
        s8,
        "Forward Chaining derives what logically follows from the facts and rules; it does not predict. "
        "Notice how count tables ensure linear time execution, and our derivation trace records exact provenance for the proof tree."
    )

    # =========================================================================
    # SLIDE 9: A* Typosquatting Detection
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "A* Heuristic Search for Typosquatting", "SECONDARY AI SEARCH ENGINE")

    add_card(s9, 0.8, 1.6, 5.7, 4.4, "A* FORMULATION: f(n) = g(n) + h(n)")
    tb = s9.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(5.3), Inches(3.6))
    tf = tb.text_frame
    tf.word_wrap = True
    astar_desc = [
        "Goal: Find minimal-cost edit path between candidate hostname label and known brand.",
        "State Space: n = (i, j) where i = source pos, j = target pos.",
        "Accumulated Cost g(n): Integer sum of operations.",
        "Heuristic Function: h(i, j) = |(|source| - i) - (|target| - j)|",
        "",
        "Operations & Integer Costs:",
        "    • MATCH: Cost 0        • SUBSTITUTE: Cost 1",
        "    • INSERT: Cost 1       • DELETE: Cost 1",
        "    • TRANSPOSITION (Damerau swap): Cost 1",
        "",
        "Admissibility Proof: Each operation reduces length discrepancy by at most 1 unit; therefore h(n) <= h*(n) always.",
    ]
    for i, a in enumerate(astar_desc):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = a
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_PRIMARY if not a.startswith("Admissibility") else PEACH_DARK

    add_card(s9, 6.8, 1.6, 5.7, 4.4, "TYPOSQUATTING EXECUTION EXAMPLE")
    tb2 = s9.shapes.add_textbox(Inches(7.0), Inches(2.2), Inches(5.3), Inches(3.6))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    ex_astar = [
        "Example Input: g00gle.com",
        "Search: 'g00gle' -> 'google'",
        "",
        "A* Path Generated:",
        "1. MATCH 'g'",
        "2. SUBSTITUTE '0' -> 'o' (Cost +1)",
        "3. SUBSTITUTE '0' -> 'o' (Cost +1)",
        "4. MATCH 'g'",
        "5. MATCH 'l'",
        "6. MATCH 'e'",
        "Total Optimal Cost: g(n) = 2",
        "",
        "Since g(n) <= 2: Asserts 'brand_typosquatting_detected'.",
    ]
    for i, e in enumerate(ex_astar):
        p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
        p.text = e
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_PRIMARY if not "Asserts" in e else EMERALD

    add_callout(s9, 0.8, 6.2, 11.7, 0.7, "CRITICAL: A* does NOT produce HIGH RISK; it produces a ground fact. Forward Chaining reasons about it.")
    set_speaker_notes(
        s9,
        "A* is our secondary AI algorithm. It solves the string edit distance problem between the hostname label and our brand lexicon. "
        "Notice that our heuristic is admissible because no operation costing 1 can reduce the length difference by more than 1. "
        "And remember: A* never assigns HIGH RISK; it only asserts brand_typosquatting_detected."
    )

    # =========================================================================
    # SLIDE 10: Explainability + QR
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)
    add_header(s10, "Explainability (Proof Trees) & QR Integration", "INTERPRETABILITY & ADAPTERS")

    add_card(s10, 0.8, 1.6, 5.7, 4.4, "PROOF TREE: https://g00gle.com/login")
    tb = s10.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(5.3), Inches(3.6))
    tf = tb.text_frame
    tf.word_wrap = True
    proof_txt = [
        "pattern_targeted_brand_phishing [CRITICAL]",
        "  └── via RULE_TYPOSQUATTING_SENSITIVE_KEYWORD",
        "       ├── brand_typosquatting_detected [GROUND FACT]",
        "       │    └── A* Search: 'g00gle' -> 'google'",
        "       │    └── Total Edit Cost = 2",
        "       │    └── SUBSTITUTE '0'->'o', SUBSTITUTE '0'->'o'",
        "       └── path_contains_sensitive_keyword [GROUND FACT]",
        "            └── Matched keyword: 'login'",
        "",
        "Resolution: Highest Severity CRITICAL  =>  Verdict: HIGH RISK",
    ]
    for i, pt in enumerate(proof_txt):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = pt
        p.font.size = Pt(10)
        p.font.color.rgb = ROSE if "CRITICAL" in pt else TEXT_PRIMARY

    add_card(s10, 6.8, 1.6, 5.7, 4.4, "OFFLINE QR CODE INTEGRATION")
    tb2 = s10.shapes.add_textbox(Inches(7.0), Inches(2.2), Inches(5.3), Inches(3.6))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    qr_txt = [
        "QR Processing Pipeline:",
        "1. Offline Image Ingestion (PNG, JPG, WebP)",
        "2. OpenCV QRCodeDetector isolates finder patterns",
        "3. Decodes embedded URL payload completely offline",
        "4. Non-URL payloads (vCards, plain text) validated & rejected",
        "5. Decoded URL passed into analyze_url()",
        "",
        "Key Design Principle:",
        "QR decoding changes the input method, NOT the security-analysis logic.",
        "Guarantees 100% semantic parity with direct URL entry.",
    ]
    for i, q in enumerate(qr_txt):
        p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
        p.text = q
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_PRIMARY if not "Key Design" in q else PEACH_DARK

    add_callout(s10, 0.8, 6.2, 11.7, 0.7, "The Proof Tree provides complete transparent provenance; every deduction is traceable to observable lexical facts.")
    set_speaker_notes(
        s10,
        "Here is the proof tree synthesized for g00gle.com/login. Every node answers 'Why?'. Why targeted brand phishing? "
        "Because of Rule R4, which required brand typosquatting from A* and a sensitive keyword in the path. "
        "And on the right, notice our QR integration is an input adapter; it does not implement separate security logic."
    )

    # =========================================================================
    # SLIDE 11: Results & Testing
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11)
    add_header(s11, "Validation & Benchmark Test Suite", "TEST RESULTS")

    add_card(s11, 0.8, 1.6, 4.0, 4.4, "FINAL TEST SUITE")
    tb = s11.shapes.add_textbox(Inches(1.0), Inches(2.1), Inches(3.6), Inches(3.7))
    tf = tb.text_frame
    tf.word_wrap = True
    test_rows = [
        "130 / 130 TESTS PASSED",
        "• URL Parser: 13 tests",
        "• Fact Generator: 18 tests",
        "• Knowledge Base: 10 tests",
        "• Forward Chaining: 13 tests",
        "• Risk Evaluator: 18 tests",
        "• Proof Tree Synthesizer: 7 tests",
        "• A* Heuristic Search: 16 tests",
        "• Application / Web UI: 12 tests",
        "• Offline QR Decoder: 10 tests",
        "• E2E Benchmark Matrix: 13 tests",
    ]
    for i, tr in enumerate(test_rows):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = tr
        p.font.size = Pt(11)
        p.font.color.rgb = EMERALD if i == 0 else TEXT_PRIMARY

    add_card(s11, 5.1, 1.6, 7.4, 4.4, "BENCHMARK MATRIX RESULTS")
    bench_table_shape = s11.shapes.add_table(8, 3, Inches(5.3), Inches(2.2), Inches(7.0), Inches(3.5))
    btable = bench_table_shape.table
    btable.columns[0].width = Inches(3.8)
    btable.columns[1].width = Inches(1.6)
    btable.columns[2].width = Inches(1.6)

    bheaders = ["Benchmark Input URL", "Verdict", "Highest Severity"]
    for c_idx, bh in enumerate(bheaders):
        c = btable.cell(0, c_idx)
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_SUBTLE
        p = c.text_frame.paragraphs[0]
        p.text = bh
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = PEACH_DARK

    bdata = [
        ("https://example.com", "SAFE", "NONE"),
        ("https://192.168.1.10/login", "HIGH RISK", "CRITICAL"),
        ("https://google.com@attacker.com/login", "HIGH RISK", "CRITICAL"),
        ("http://example.com:8080/browse", "REVIEW", "MEDIUM"),
        ("https://g00gle.com/login", "HIGH RISK", "CRITICAL"),
        ("http://paypa1.com/index", "SUSPICIOUS", "HIGH"),
        ("QR -> https://g00gle.com/login", "HIGH RISK", "CRITICAL"),
    ]
    for r_idx, r_val in enumerate(bdata, start=1):
        for c_idx, c_val in enumerate(r_val):
            c = btable.cell(r_idx, c_idx)
            c.fill.solid()
            c.fill.fore_color.rgb = CARD_BG
            p = c.text_frame.paragraphs[0]
            p.text = c_val
            p.font.size = Pt(10)
            p.font.color.rgb = EMERALD if c_val == "SAFE" else (ROSE if c_val == "HIGH RISK" else (AMBER if c_val == "SUSPICIOUS" else TEXT_PRIMARY))

    add_callout(s11, 0.8, 6.2, 11.7, 0.7, "Zero Network Violations: Automated autouse socket monkeypatching verifies 100% offline execution.")
    set_speaker_notes(
        s11,
        "We have 130 passing tests executed in 1.1 seconds. Every test runs with an autouse socket monkeypatch that guarantees "
        "zero network calls. The benchmark table shows our 7 core test scenarios, covering clean links, direct IPs, userinfo masking, "
        "A* typosquatting with HTTPS vs HTTP, non-standard ports, and offline QR decoding."
    )

    # =========================================================================
    # SLIDE 12: Conclusion & Future Scope
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_background(s12)
    add_header(s12, "Conclusion & Future Work", "SUMMARY & WRAP-UP")

    add_card(s12, 0.8, 1.6, 5.7, 4.4, "ACCOMPLISHMENTS & LIMITATIONS")
    tb = s12.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(5.3), Inches(3.6))
    tf = tb.text_frame
    tf.word_wrap = True
    body_s12_1 = [
        "Key Accomplishments:",
        "• Pure Classical AI implementation: Horn Logic + Forward Chaining.",
        "• Upstream A* search with mathematically proven admissible heuristic.",
        "• Complete causal proof tree generation answering 'Why?'.",
        "• Zero network dependency, zero black-box ML/DL/LLMs.",
        "",
        "Transparent Limitations:",
        "• Air-gapped; cannot inspect dynamically served server payloads.",
        "• No live threat intelligence feeds or DNS resolution.",
        "• Upstream typosquatting bounded by local brand lexicon.",
        "• SAFE verdict indicates absence of local suspicious patterns.",
    ]
    for i, b in enumerate(body_s12_1):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = b
        p.font.size = Pt(11)
        p.font.color.rgb = PEACH_DARK if b.startswith("Key") or b.startswith("Transparent") else TEXT_PRIMARY

    add_card(s12, 6.8, 1.6, 5.7, 4.4, "FUTURE SCOPE")
    tb2 = s12.shapes.add_textbox(Inches(7.0), Inches(2.2), Inches(5.3), Inches(3.6))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    body_s12_2 = [
        "1. Dynamic Knowledge Base Expansion:",
        "    • Support user-authored Horn clauses via declarative JSON schema.",
        "2. Multi-Region Brand Lexicons:",
        "    • Expand local brand dictionaries for regional institutions.",
        "3. Advanced Subdomain & Homoglyph Detection:",
        "    • Add symbolic rules for internationalized punycode homoglyphs.",
        "4. Enhanced QR Image Preprocessing:",
        "    • Adaptive local contrast filtering for degraded QR codes.",
        "5. Auditable Security Logs:",
        "    • Export proof trees as standardized cryptographic audit trails.",
    ]
    for i, b in enumerate(body_s12_2):
        p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
        p.text = b
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_PRIMARY

    add_callout(s12, 0.8, 6.2, 11.7, 0.7, '"The goal is not to replace real-world threat intelligence, but to demonstrate how Classical AI can perform transparent, deterministic and explainable security reasoning."')
    set_speaker_notes(
        s12,
        "The goal is not to replace real-world threat intelligence, but to demonstrate how Classical AI can perform "
        "transparent, deterministic and explainable security reasoning. Thank you, and we welcome your questions!"
    )

    prs.save("presentation.pptx")
    print("Saved minimalist peach presentation.pptx successfully!")


if __name__ == "__main__":
    create_deck()
